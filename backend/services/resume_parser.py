import re
import io
from typing import Dict, Any, List

# Technical taxonomy for comprehensive skill extraction
TECH_KEYWORDS = {
    # Languages
    "python", "javascript", "typescript", "java", "c++", "c#", "c", "go", "golang", 
    "rust", "ruby", "php", "swift", "kotlin", "scala", "r", "sql", "bash", "shell",
    # Frontend
    "react", "react.js", "next.js", "vue", "vue.js", "angular", "svelte", "html", 
    "html5", "css", "css3", "tailwind", "tailwind css", "bootstrap", "sass", "redux",
    # Backend & Frameworks
    "fastapi", "django", "flask", "express", "express.js", "node.js", "nodejs", 
    "spring", "spring boot", "ruby on rails", "asp.net", "graphql", "rest", "restful api",
    # Databases & Storage
    "postgresql", "postgres", "mysql", "sqlite", "mongodb", "redis", "elasticsearch", 
    "cassandra", "dynamodb", "supabase", "firebase", "neo4j",
    # Cloud & DevOps
    "aws", "amazon web services", "azure", "gcp", "google cloud", "docker", 
    "kubernetes", "k8s", "terraform", "ci/cd", "github actions", "gitlab ci", 
    "jenkins", "linux", "nginx", "ansible",
    # AI / ML & Data
    "machine learning", "deep learning", "nlp", "llm", "rag", "pytorch", "tensorflow", 
    "keras", "scikit-learn", "pandas", "numpy", "opencv", "hugging face", "transformers",
    # Tools & Practices
    "git", "github", "gitlab", "jira", "agile", "scrum", "unit testing", "pytest", 
    "jest", "cypress", "postman", "figma"
}

def extract_text_from_pdf(file_bytes: bytes) -> str:
    text = ""
    try:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
    except Exception as e:
        # Fallback to simple binary decode or basic extraction
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            for p in reader.pages:
                t = p.extract_text()
                if t:
                    text += t + "\n"
        except Exception:
            pass
    return text.strip()

def extract_text_from_docx(file_bytes: bytes) -> str:
    text = ""
    try:
        import docx
        doc = docx.Document(io.BytesIO(file_bytes))
        for p in doc.paragraphs:
            if p.text:
                text += p.text + "\n"
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        text += cell.text + " "
                text += "\n"
    except Exception as e:
        pass
    return text.strip()

def extract_skills(text: str) -> List[str]:
    found = set()
    lower_text = " " + re.sub(r"[^\w\s\+\#\/\.\-]", " ", text.lower()) + " "
    
    for kw in TECH_KEYWORDS:
        # Exact word/token boundary match
        pattern = r"(?<!\w)" + re.escape(kw) + r"(?!\w)"
        if re.search(pattern, lower_text):
            found.add(kw.title() if len(kw) > 3 else kw.upper())
            
    return sorted(list(found))

def is_section_header(line: str, keywords: List[str]) -> bool:
    clean = line.strip().lower()
    if len(clean) > 45 or clean.startswith(("-", "•", "*", "–")):
        return False
    clean_alpha = re.sub(r"[^a-z\s]", "", clean).strip()
    return any(clean_alpha == kw or clean_alpha.startswith(kw) or clean_alpha.endswith(kw) for kw in keywords)

def extract_sections(text: str) -> Dict[str, Any]:
    lines = [l.strip() for l in text.split("\n") if l.strip()]
    
    skills = extract_skills(text)
    
    # Categorize tools vs skills
    tools_list = [s for s in skills if s.lower() in {
        "git", "github", "gitlab", "docker", "kubernetes", "k8s", "linux", 
        "jira", "postman", "figma", "jenkins", "terraform", "pytest", "jest"
    }]
    
    # Experience bullets extraction
    experience_bullets = []
    projects = []
    education = []
    
    current_section = "summary"
    for line in lines:
        if is_section_header(line, ["work experience", "professional experience", "experience", "employment"]):
            current_section = "experience"
            continue
        elif is_section_header(line, ["selected projects", "projects", "personal projects", "academic projects"]):
            current_section = "projects"
            continue
        elif is_section_header(line, ["education", "academics", "university", "degrees"]):
            current_section = "education"
            continue
        elif is_section_header(line, ["technical skills", "skills", "technologies", "core competencies"]):
            current_section = "skills"
            continue

        if current_section == "experience":
            if line.startswith(("-", "•", "*", "–")) or len(line) > 30:
                clean_bullet = line.lstrip("-•*– ").strip()
                if len(clean_bullet) > 15:
                    experience_bullets.append(clean_bullet)
        elif current_section == "projects":
            if line.startswith(("-", "•", "*", "–")) or len(line) > 20:
                projects.append(line.lstrip("-•*– ").strip())
        elif current_section == "education":
            if any(deg in lower_line for deg in ["bachelor", "master", "phd", "b.s.", "m.s.", "b.tech", "degree", "university", "college", "gpa"]):
                education.append(line)

    return {
        "skills": skills,
        "tools": tools_list,
        "experience": experience_bullets[:20],
        "projects": projects[:15],
        "education": education[:5]
    }

def parse_resume(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    filename_lower = filename.lower()
    if filename_lower.endswith(".pdf"):
        raw_text = extract_text_from_pdf(file_bytes)
    elif filename_lower.endswith(".docx") or filename_lower.endswith(".doc"):
        raw_text = extract_text_from_docx(file_bytes)
    else:
        # Fallback to UTF-8 decode
        try:
            raw_text = file_bytes.decode("utf-8")
        except Exception:
            raw_text = file_bytes.decode("latin-1", errors="ignore")

    structured = extract_sections(raw_text)
    return {
        "raw_text": raw_text,
        "parsed_json": structured
    }
