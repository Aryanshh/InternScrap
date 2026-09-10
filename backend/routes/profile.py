import os
import io
import logging
from datetime import datetime
from typing import Optional, List, Dict, Any
from fastapi import APIRouter, Depends, HTTPException, Response, Header, Query
from pydantic import BaseModel
from sqlalchemy.orm import Session
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from backend.database import get_db
from backend.models.profile import UserProfile
from backend.models.resume import Resume

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api/profile", tags=["Profile"])

class ProfileUpdateRequest(BaseModel):
    full_name: Optional[str] = None
    email: Optional[str] = None
    phone: Optional[str] = None
    location: Optional[str] = None
    headline: Optional[str] = None
    bio: Optional[str] = None
    github_url: Optional[str] = None
    linkedin_url: Optional[str] = None
    portfolio_url: Optional[str] = None
    desired_work_mode: Optional[str] = None
    min_salary: Optional[int] = None
    min_hourly_rate: Optional[float] = None
    target_platforms: Optional[List[str]] = None
    skills: Optional[List[str]] = None
    experience: Optional[List[Dict[str, Any]]] = None
    education: Optional[List[Dict[str, Any]]] = None
    digest_enabled: Optional[bool] = None
    digest_frequency: Optional[str] = None
    digest_min_score: Optional[float] = None

class SwitchProfileRequest(BaseModel):
    user_id: str

DEFAULT_PROFILES: Dict[str, Dict[str, Any]] = {
    "demo": {
        "id": "demo",
        "full_name": "Demo Candidate",
        "email": "demo@internscrap.dev",
        "phone": "+1 (555) 123-4567",
        "location": "Austin, TX / Remote",
        "headline": "Junior Full-Stack Developer & Technical Intern",
        "bio": (
            "Passionate early-career developer focused on building interactive web applications with modern React, TypeScript, "
            "and Python. Seeking high-impact software engineering internships, fellowships, and junior developer opportunities."
        ),
        "github_url": "https://github.com/demo",
        "linkedin_url": "https://linkedin.com/in/demo-intern",
        "portfolio_url": "https://demo.dev",
        "desired_work_mode": "remote",
        "min_salary": 75000,
        "min_hourly_rate": 35.0,
        "target_platforms": ["Wellfound", "Outlier", "Remotive", "Arbeitnow"],
        "skills": [
            "JavaScript", "TypeScript", "React", "Python", "HTML5", "CSS3", "Git",
            "SQL", "Tailwind CSS", "REST APIs", "FastAPI", "Express"
        ],
        "experience": [
            {
                "company": "Campus Tech Labs",
                "role": "Software Developer Intern",
                "location": "Remote",
                "start_date": "May 2024",
                "end_date": "Aug 2024",
                "bullets": [
                    "Built dynamic React UI components with Tailwind CSS for internal student portal used by 2,000+ active campus members.",
                    "Created RESTful endpoints in Python FastAPI with SQLite to handle course scheduling queries with sub-second response times.",
                    "Authored automated unit tests ensuring 90%+ code coverage before production deployment."
                ]
            }
        ],
        "education": [
            {
                "institution": "University of Texas, Austin",
                "degree": "B.S. in Computer Science",
                "grad_year": "2026",
                "gpa": "3.75 / 4.0"
            }
        ],
        "digest_enabled": True,
        "digest_frequency": "12h",
        "digest_min_score": 60.0
    },
    "Aryanshh": {
        "id": "Aryanshh",
        "full_name": "Aryan Sharma",
        "email": "",
        "phone": "",
        "location": "",
        "headline": "",
        "bio": "",
        "github_url": "",
        "linkedin_url": "",
        "portfolio_url": "",
        "desired_work_mode": "remote",
        "min_salary": 0,
        "min_hourly_rate": 0.0,
        "target_platforms": [],
        "skills": [],
        "experience": [],
        "education": [],
        "digest_enabled": False,
        "digest_frequency": "daily",
        "digest_min_score": 70.0
    },
    "Nishtha": {
        "id": "Nishtha",
        "full_name": "Nishtha",
        "email": "",
        "phone": "",
        "location": "",
        "headline": "",
        "bio": "",
        "github_url": "",
        "linkedin_url": "",
        "portfolio_url": "",
        "desired_work_mode": "remote",
        "min_salary": 0,
        "min_hourly_rate": 0.0,
        "target_platforms": [],
        "skills": [],
        "experience": [],
        "education": [],
        "digest_enabled": False,
        "digest_frequency": "daily",
        "digest_min_score": 70.0
    }
}

def seed_all_profiles(db: Session, force_reset_blanks: bool = True):
    """Ensure demo has mock data, and Aryanshh & Nishtha are reset to clean blank state for self-setup."""
    for pid, data in DEFAULT_PROFILES.items():
        existing = db.query(UserProfile).filter(UserProfile.id == pid).first()
        if not existing:
            new_p = UserProfile(**data)
            db.add(new_p)
        elif pid in ("Aryanshh", "Nishtha") and force_reset_blanks:
            # Check if existing profile contains old mock data (e.g. TechNova or Apex AI) and wipe clean
            has_old_mock = False
            for exp in (existing.experience or []):
                if exp.get("company") in ("TechNova Solutions", "Apex AI Labs", "Cognitive Insights", "DataSphere AI"):
                    has_old_mock = True
                    break
            if has_old_mock:
                for k, v in data.items():
                    if k != "id":
                        setattr(existing, k, v)
    db.commit()

def resolve_target_user_id(user_id: Optional[str], x_user_id: Optional[str]) -> str:
    chosen = (user_id or x_user_id or "Aryanshh").strip()
    if not chosen or chosen == "default_user":
        chosen = "Aryanshh"
    return chosen

def get_or_create_profile(db: Session, user_id: Optional[str] = None) -> UserProfile:
    target_id = resolve_target_user_id(user_id, None)
    profile = db.query(UserProfile).filter(UserProfile.id == target_id).first()
    if not profile:
        seed_data = DEFAULT_PROFILES.get(target_id)
        if seed_data:
            profile = UserProfile(**seed_data)
        else:
            profile = UserProfile(id=target_id, full_name=target_id)
        db.add(profile)
        db.commit()
        db.refresh(profile)
    return profile

@router.get("/list")
def list_profiles(db: Session = Depends(get_db)):
    """Return the three candidate login profiles: demo, Aryanshh, and Nishtha."""
    seed_all_profiles(db)
    profiles = db.query(UserProfile).all()
    
    order = ["demo", "Aryanshh", "Nishtha"]
    allowed = {"demo", "Aryanshh", "Nishtha"}
    filtered = [p for p in profiles if p.id in allowed]
    ordered = sorted(filtered, key=lambda p: order.index(p.id))
    
    return [
        {
            "id": p.id,
            "full_name": p.full_name,
            "headline": p.headline if p.headline else "Profile Setup Pending (Self-Configured)",
            "email": p.email or "",
            "avatar": p.full_name[:1].upper() if p.full_name else p.id[:1].upper(),
            "role_tag": (
                "Self-Setup" if p.id in ("Aryanshh", "Nishtha")
                else "Demo / Intern" if p.id == "demo"
                else "Candidate"
            ),
            "skills_count": len(p.skills or []),
            "desired_work_mode": p.desired_work_mode,
            "min_hourly_rate": p.min_hourly_rate,
            "target_platforms": p.target_platforms or [],
        }
        for p in ordered
    ]

@router.post("/switch")
def switch_profile(payload: SwitchProfileRequest, db: Session = Depends(get_db)):
    """Switch active profile session."""
    target_id = payload.user_id.strip()
    profile = get_or_create_profile(db, user_id=target_id)
    return {
        "message": f"Switched to profile: {profile.full_name} ({profile.id})",
        "active_user_id": profile.id,
        "profile": {
            "id": profile.id,
            "full_name": profile.full_name,
            "headline": profile.headline,
            "email": profile.email,
        }
    }

@router.get("")
def get_profile(
    user_id: Optional[str] = Query(None),
    x_user_id: Optional[str] = Header(None, alias="X-User-Id"),
    db: Session = Depends(get_db)
):
    """Retrieve candidate profile with active resume status."""
    target_id = resolve_target_user_id(user_id, x_user_id)
    profile = get_or_create_profile(db, target_id)
    active_resume = db.query(Resume).filter(Resume.is_active == True).first()

    return {
        "id": profile.id,
        "full_name": profile.full_name,
        "email": profile.email,
        "phone": profile.phone,
        "location": profile.location,
        "headline": profile.headline,
        "bio": profile.bio,
        "github_url": profile.github_url,
        "linkedin_url": profile.linkedin_url,
        "portfolio_url": profile.portfolio_url,
        "desired_work_mode": profile.desired_work_mode,
        "min_salary": profile.min_salary,
        "min_hourly_rate": profile.min_hourly_rate,
        "target_platforms": profile.target_platforms or [],
        "skills": profile.skills or [],
        "experience": profile.experience or [],
        "education": profile.education or [],
        "digest_enabled": profile.digest_enabled,
        "digest_frequency": profile.digest_frequency,
        "digest_min_score": profile.digest_min_score,
        "updated_at": profile.updated_at.isoformat() if profile.updated_at else None,
        "active_resume": {
            "id": active_resume.id,
            "filename": active_resume.filename,
            "uploaded_at": active_resume.created_at.isoformat() if active_resume.created_at else None,
        } if active_resume else None,
    }

@router.put("")
def update_profile(
    data: ProfileUpdateRequest,
    user_id: Optional[str] = Query(None),
    x_user_id: Optional[str] = Header(None, alias="X-User-Id"),
    db: Session = Depends(get_db)
):
    """Update profile preferences, skills, and work history."""
    target_id = resolve_target_user_id(user_id, x_user_id)
    profile = get_or_create_profile(db, target_id)

    update_dict = data.model_dump(exclude_unset=True)
    for key, value in update_dict.items():
        setattr(profile, key, value)

    profile.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(profile)

    return {
        "message": "Profile updated successfully",
        "profile": {
            "id": profile.id,
            "full_name": profile.full_name,
            "email": profile.email,
            "headline": profile.headline,
            "skills_count": len(profile.skills or []),
            "updated_at": profile.updated_at.isoformat(),
        }
    }

@router.post("/sync-from-resume")
def sync_from_active_resume(
    user_id: Optional[str] = Query(None),
    x_user_id: Optional[str] = Header(None, alias="X-User-Id"),
    db: Session = Depends(get_db)
):
    """Extract skills and experience from the active uploaded resume and merge into targeted profile."""
    target_id = resolve_target_user_id(user_id, x_user_id)
    profile = get_or_create_profile(db, target_id)
    active_resume = db.query(Resume).filter(Resume.is_active == True).first()

    if not active_resume or not active_resume.parsed_json:
        raise HTTPException(status_code=400, detail="No active parsed resume found to sync from.")

    parsed = active_resume.parsed_json
    synced_items = []

    # 1. Sync skills
    if "skills" in parsed and isinstance(parsed["skills"], list):
        current_skills = set(profile.skills or [])
        new_skills = [s for s in parsed["skills"] if s not in current_skills]
        profile.skills = list(current_skills.union(set(parsed["skills"])))
        synced_items.append(f"{len(new_skills)} new skills")

    # 2. Sync experience bullets if empty
    exp_bullets = parsed.get("experience") or parsed.get("experience_bullets") or []
    if exp_bullets and isinstance(exp_bullets, list):
        if not profile.experience:
            profile.experience = [
                {
                    "company": "Professional Experience",
                    "role": profile.headline or "Software Contributor",
                    "location": "Remote",
                    "start_date": "2023",
                    "end_date": "Present",
                    "bullets": exp_bullets[:4],
                }
            ]
            synced_items.append("experience bullets")

    # 3. Sync education if available
    if "education" in parsed and isinstance(parsed["education"], list) and parsed["education"]:
        edu_list = []
        for e in parsed["education"]:
            if isinstance(e, dict):
                edu_list.append({
                    "institution": e.get("institution", "University"),
                    "degree": e.get("degree", "Degree"),
                    "grad_year": e.get("grad_year") or e.get("year", "2024"),
                    "gpa": e.get("gpa", ""),
                })
            elif isinstance(e, str):
                edu_list.append({
                    "institution": e,
                    "degree": "Degree / Study",
                    "grad_year": "2024",
                    "gpa": "",
                })
        if edu_list:
            profile.education = edu_list
            synced_items.append("education")

    profile.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(profile)

    return {
        "message": f"Successfully synced from resume for {profile.full_name}: {', '.join(synced_items) if synced_items else 'Profile already up to date.'}",
        "skills": profile.skills,
        "experience": profile.experience,
        "education": profile.education,
    }

@router.get("/export-docx")
def export_profile_docx(
    user_id: Optional[str] = Query(None),
    x_user_id: Optional[str] = Header(None, alias="X-User-Id"),
    db: Session = Depends(get_db)
):
    """Compile the user profile into an ATS-formatted DOCX resume."""
    target_id = resolve_target_user_id(user_id, x_user_id)
    profile = get_or_create_profile(db, target_id)

    doc = Document()

    # Margins: 0.75 in
    for sec in doc.sections:
        sec.top_margin = Inches(0.75)
        sec.bottom_margin = Inches(0.75)
        sec.left_margin = Inches(0.75)
        sec.right_margin = Inches(0.75)

    NAVY = RGBColor(0x1E, 0x3A, 0x8A)
    CHARCOAL = RGBColor(0x1E, 0x29, 0x3B)
    SLATE = RGBColor(0x47, 0x55, 0x69)

    # Name Header
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_name = p_name.add_run(profile.full_name or "Candidate")
    run_name.bold = True
    run_name.font.size = Pt(20)
    run_name.font.color.rgb = NAVY

    # Contact Line
    p_contact = doc.add_paragraph()
    p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_parts = [p for p in [profile.email, profile.phone, profile.location, profile.github_url, profile.linkedin_url] if p]
    run_contact = p_contact.add_run(" | ".join(contact_parts))
    run_contact.font.size = Pt(9.5)
    run_contact.font.color.rgb = SLATE

    def add_sec_heading(title: str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(title.upper())
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = NAVY

    # Summary
    if profile.bio:
        add_sec_heading("Professional Summary")
        p_bio = doc.add_paragraph()
        p_bio.paragraph_format.space_after = Pt(4)
        run_bio = p_bio.add_run(profile.bio)
        run_bio.font.size = Pt(10)
        run_bio.font.color.rgb = CHARCOAL

    # Technical Skills
    if profile.skills:
        add_sec_heading("Technical Proficiencies")
        p_sk = doc.add_paragraph()
        p_sk.paragraph_format.space_after = Pt(4)
        run_lbl = p_sk.add_run("Core Skills: ")
        run_lbl.bold = True
        run_lbl.font.size = Pt(10)
        run_lbl.font.color.rgb = CHARCOAL
        run_sk = p_sk.add_run(", ".join(profile.skills))
        run_sk.font.size = Pt(10)
        run_sk.font.color.rgb = CHARCOAL

    # Work Experience
    if profile.experience:
        add_sec_heading("Professional Experience")
        for exp in profile.experience:
            p_exp = doc.add_paragraph()
            p_exp.paragraph_format.space_before = Pt(4)
            p_exp.paragraph_format.space_after = Pt(2)
            r_role = p_exp.add_run(exp.get("role", "Role"))
            r_role.bold = True
            r_role.font.size = Pt(10.5)
            r_role.font.color.rgb = CHARCOAL

            r_co = p_exp.add_run(f" | {exp.get('company', '')}")
            r_co.font.size = Pt(10.5)
            r_co.font.color.rgb = NAVY

            dates = f"{exp.get('start_date', '')} - {exp.get('end_date', '')}".strip(" -")
            if dates:
                r_dates = p_exp.add_run(f" ({dates})")
                r_dates.italic = True
                r_dates.font.size = Pt(9.5)
                r_dates.font.color.rgb = SLATE

            for bullet in exp.get("bullets", []):
                p_b = doc.add_paragraph(style="List Bullet")
                p_b.paragraph_format.space_after = Pt(2)
                p_b.paragraph_format.line_spacing = 1.15
                r_b = p_b.add_run(bullet)
                r_b.font.size = Pt(10)
                r_b.font.color.rgb = CHARCOAL

    # Education
    if profile.education:
        add_sec_heading("Education")
        for edu in profile.education:
            p_edu = doc.add_paragraph()
            p_edu.paragraph_format.space_after = Pt(2)
            r_deg = p_edu.add_run(edu.get("degree", "Degree"))
            r_deg.bold = True
            r_deg.font.size = Pt(10)
            r_inst = p_edu.add_run(f" - {edu.get('institution', '')}")
            r_inst.font.size = Pt(10)
            extra = [edu.get("grad_year", ""), edu.get("gpa", "")]
            extra_str = ", ".join([x for x in extra if x])
            if extra_str:
                r_extra = p_edu.add_run(f" ({extra_str})")
                r_extra.italic = True
                r_extra.font.size = Pt(9.5)
                r_extra.font.color.rgb = SLATE

    buffer = io.BytesIO()
    doc.save(buffer)
    file_bytes = buffer.getvalue()

    safe_name = (profile.full_name or profile.id or "Candidate").replace(" ", "_")
    filename = f"{safe_name}_Profile_Resume.docx"
    return Response(
        content=file_bytes,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
